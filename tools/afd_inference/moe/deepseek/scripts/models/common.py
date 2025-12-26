import os
import time
import math
import numpy as np
import torch
from torch import nn
import torch.nn.functional as F
import torch_npu
import torch.distributed as dist
from torch.distributed.distributed_c10d import _world
import torchair as tng

from transformers.utils import (
    add_start_docstrings,
    add_start_docstrings_to_model_forward,
    is_flash_attn_2_available,
    is_flash_attn_greater_or_equal_2_10,
    logging,
    replace_return_docstrings,
)
from transformers.modeling_attn_mask_utils import (
    AttentionMaskConverter,
    _prepare_4d_attention_mask,
    _prepare_4d_causal_attention_mask,
)
from transformers.utils.import_utils import is_torch_fx_available
from transformers.pytorch_utils import (
    ALL_LAYERNORM_LAYERS,
    is_torch_greater_or_equal_than_1_13,
)
from transformers.modeling_utils import PreTrainedModel

if is_flash_attn_2_available():
    from flash_attn import flash_attn_func, flash_attn_varlen_func
    from flash_attn.bert_padding import index_first_axis, pad_input, unpad_input  # noqa

# This makes `_prepare_4d_causal_attention_mask` a leaf function in the FX graph.
# It means that the function will not be traced through and simply appear as a node in the graph.
if is_torch_fx_available():
    if not is_torch_greater_or_equal_than_1_13:
        import torch.fx

    _prepare_4d_causal_attention_mask = torch.fx.wrap(_prepare_4d_causal_attention_mask)

from .global_setting import _TODO_REQUIRE_API, _PAGE_ATTENTION_SETTING
from .configuration_deepseek import DeepseekV2Config

logging.set_verbosity_info()
logger = logging.get_logger("transformers")

dtype_mapping = {
    "float": torch.float,
    "float16": torch.float16,
    "bfloat16": torch.bfloat16,
}
set_dtype = os.getenv("DTYPE", "bfloat16")
model_dtype = dtype_mapping.get(set_dtype, torch.bfloat16)
quant_dtype_mapping = {
    "float": torch.float,
    "float16": torch.float,
    "bfloat16": torch.bfloat16,
}
# quant_dtype = quant_dtype_mapping.get(set_dtype, torch.bfloat16)
quant_dtype = torch.float
exe_mode = os.getenv("EXE_MODE")
ffn_mode = os.getenv("FFN_MODE")
save_model_switch = int(os.getenv("SAVE_MODEL", "0"))
actual_seq_len = int(os.getenv("ACTUAL_SEQ_LEN", "256"))
max_new_tokens = int(os.getenv("MAX_NEW_TOKENS", "32"))


def fix_rand_seed(seed=42):
    torch.manual_seed(seed)
    torch.npu.manual_seed_all(seed)


def sync_and_get_time(start_time=None, model_name="dsV3", use_syn=True):
    if use_syn:
        torch.npu.synchronize()
    time_stamp = time.time()
    if start_time is not None:
        time_stamp -= start_time
        logger.info(f"{model_name} inference time cost is: {time_stamp*1000:.2f} ms" )
    return time_stamp


def process_run_time(run_time_list):
    on_cloud = int(os.getenv("ON_CLOUD", "0"))
    local_steps = min(int(os.getenv("MAX_NEW_TOKENS", "50")), 20)
    run_steps = 100 if on_cloud else local_steps

    last_run_time_list = run_time_list[-run_steps:]
    mean_time = np.mean(last_run_time_list)
    refine_run_time_list = []
    for time in last_run_time_list:
        if time >= mean_time * 3:
            continue
        refine_run_time_list.append(time)
    if len(refine_run_time_list) != len(last_run_time_list):
        flag = "abnormal"
    else:
        flag = "normal"
    return np.mean(refine_run_time_list), flag


def apply_quant(x):
    amax, _ = torch.max(torch.abs(x), -1, keepdim=True)
    scale = amax / 127.0
    x = (x / scale).to(torch.int8)
    return x


def _get_unpad_data(attention_mask):
    seqlens_in_batch = attention_mask.sum(dim=-1, dtype=torch.int32)
    indices = torch.nonzero(attention_mask.flatten(), as_tuple=False).flatten()
    max_seqlen_in_batch = seqlens_in_batch.max().item()
    cu_seqlens = F.pad(
        torch.cumsum(seqlens_in_batch, dim=0, dtype=torch.torch.int32), (1, 0)
    )
    return (
        indices,
        cu_seqlens,
        max_seqlen_in_batch,
    )

def ep_comm(input_tensor, output_tensor=None, mode="all_gather", group=_world._default_pg, jump_flag=False):
    if jump_flag:
        return input_tensor
    assert mode in ["all_gather", "reduce_scatter", "all2all"]

    orig_shape = input_tensor.shape
    hidden_size = orig_shape[-1]
    if mode == "all2all":
        input_tensor = input_tensor.view(-1)
        dist.all_to_all_single(input_tensor, input_tensor, group=group)
        # input_tensor = self.all2all_prefill
        comm_output = input_tensor.view(*orig_shape)
    elif mode == "all_gather":
        input_tensor = input_tensor.view(-1, hidden_size)
        dist.all_gather_into_tensor(output_tensor, input_tensor, group=group)
        comm_output = output_tensor
    elif mode == "reduce_scatter": # sum reduce on dim 0
        input_tensor = input_tensor.view(-1, hidden_size)
        dist.reduce_scatter_tensor(output_tensor, input_tensor, group=group)
        comm_output = output_tensor
    else:
        raise NotImplementedError(
            f"insupportable DeepseekV2MoeInferCommMode for mode as: {mode}"
        )

    return comm_output

class DeepseekV2RMSNorm(nn.Module):
    def __init__(self, hidden_size, eps=1e-6):
        """
        DeepseekV2RMSNorm is equivalent to T5LayerNorm
        """
        super().__init__()
        self.weight = nn.Parameter(torch.ones(hidden_size, dtype=model_dtype))
        self.variance_epsilon = eps

    def ln(self, hidden_states):
        input_dtype = hidden_states.dtype
        hidden_states = hidden_states.to(torch.float32)
        variance = hidden_states.pow(2).mean(-1, keepdim=True)
        hidden_states = hidden_states * torch.rsqrt(variance + self.variance_epsilon)
        return self.weight * hidden_states.to(input_dtype)

    def ln_npu(self, hidden_states):
        result = torch_npu.npu_rms_norm(hidden_states, self.weight, self.variance_epsilon)[0]
        return result

    def forward(self, hidden_states, *args):
        if len(args) == 0: # only hidden_states exists
            result = self.ln_npu(hidden_states)
            return result
        elif len(args) == 1 and args[0] is None: # residual is None
            result = self.ln_npu(hidden_states)
            residual = hidden_states
            return (result, residual)
        elif len(args) == 1: # residual is not None:
            residual = args[0]
            y, _, x = torch_npu.npu_add_rms_norm(residual, hidden_states, self.weight, self.variance_epsilon)
            return (y, x)
        else:
            raise NotImplementedError(
                f"insupportable DeepseekV2RMSNorm for input_args len as: {len(args)+1}"
            )


class DeepseekV2RotaryEmbedding(nn.Module):
    def __init__(self, config, dim, max_position_embeddings=2048, base=10000, device=None):
        super().__init__()
        self.config = config
        self.dim = dim
        self.max_position_embeddings = max_position_embeddings
        self.base = base
        inv_freq = 1.0 / (
            self.base ** (torch.arange(0, self.dim, 2).float().to(device) / self.dim)
        )
        self.register_buffer("inv_freq", inv_freq, persistent=False)
        self.next_n = int(os.getenv("NEXT_N", "0"))
        self.spec_len = self.next_n + 1

        # Build here to make `torch.jit.trace` work.
        self._set_cos_sin_cache(
            seq_len=max_position_embeddings,
            device=self.inv_freq.device,
            dtype=torch.get_default_dtype(),
        )

    def _set_cos_sin_cache(self, seq_len, device, dtype):
        self.max_seq_len_cached = seq_len
        t = torch.arange(
            self.max_seq_len_cached, device=device, dtype=self.inv_freq.dtype
        )

        freqs = torch.outer(t, self.inv_freq.to(t.device))
        # Different from paper, but it uses a different permutation in order to obtain the same calculation
        emb = torch.cat((freqs, freqs), dim=-1)
        self.register_buffer("cos_cached", emb.cos().to(dtype), persistent=False)
        self.register_buffer("sin_cached", emb.sin().to(dtype), persistent=False)

    def __forward(self, x, kv_len=None):
        # x: [bs, num_attention_heads, seq_len, head_size]
        if self.max_seq_len_cached is None or kv_len > self.max_seq_len_cached:
            self._set_cos_sin_cache(seq_len=kv_len, device=x.device, dtype=x.dtype)

        return (
            self.cos_cached[:kv_len].to(dtype=x.dtype),
            self.sin_cached[:kv_len].to(dtype=x.dtype),
        )

    def forward(self, batch_size, seq_len, kv_len, max_seq_len=None):
        # x shape is [bs, num_attention_heads, seq_len, head_size]
        if max_seq_len is None:
            self._set_cos_sin_cache(seq_len=kv_len, device="npu", dtype=model_dtype)
        elif max_seq_len > self.max_seq_len_cached:
            self._set_cos_sin_cache(seq_len=max_seq_len, device="npu", dtype=model_dtype)

        # SD -> BNSD
        cos = self.cos_cached[:seq_len].view(1, 1, seq_len, -1).repeat(batch_size, 1, 1, 1)
        sin = self.sin_cached[:seq_len].view(1, 1, seq_len, -1).repeat(batch_size, 1, 1, 1)

        return (
            cos.to(dtype=model_dtype),
            sin.to(dtype=model_dtype),
        )


# Copied from transformers.models.llama.modeling_llama.LlamaLinearScalingRotaryEmbedding with Llama->DeepseekV2
class DeepseekV2LinearScalingRotaryEmbedding(DeepseekV2RotaryEmbedding):
    """DeepseekV2RotaryEmbedding extended with linear scaling. Credits to the Reddit user /u/kaiokendev"""

    def __init__(
        self,
        config,
        dim,
        max_position_embeddings=2048,
        base=10000,
        device=None,
        scaling_factor=1.0,
    ):
        self.scaling_factor = scaling_factor
        super().__init__(config, dim, max_position_embeddings, base, device)

    def _set_cos_sin_cache(self, seq_len, device, dtype):
        self.max_seq_len_cached = seq_len
        t = torch.arange(
            self.max_seq_len_cached, device=device, dtype=self.inv_freq.dtype
        )
        t = t / self.scaling_factor

        freqs = torch.outer(t, self.inv_freq)
        # Different from paper, but it uses a different permutation in order to obtain the same calculation
        emb = torch.cat((freqs, freqs), dim=-1)
        self.register_buffer("cos_cached", emb.cos().to(dtype), persistent=False)
        self.register_buffer("sin_cached", emb.sin().to(dtype), persistent=False)


# Copied from transformers.models.llama.modeling_llama.LlamaDynamicNTKScalingRotaryEmbedding with Llama->DeepseekV2
class DeepseekV2DynamicNTKScalingRotaryEmbedding(DeepseekV2RotaryEmbedding):
    """DeepseekV2RotaryEmbedding extended with Dynamic NTK scaling. Credits to the Reddit users /u/bloc97 and /u/emozilla"""

    def __init__(
        self,
        config,
        dim,
        max_position_embeddings=2048,
        base=10000,
        device=None,
        scaling_factor=1.0,
    ):
        self.scaling_factor = scaling_factor
        super().__init__(config, dim, max_position_embeddings, base, device)

    def _set_cos_sin_cache(self, seq_len, device, dtype):
        self.max_seq_len_cached = seq_len

        if seq_len > self.max_position_embeddings:
            base = self.base * (
                (self.scaling_factor * seq_len / self.max_position_embeddings)
                - (self.scaling_factor - 1)
            ) ** (self.dim / (self.dim - 2))
            inv_freq = 1.0 / (
                base ** (torch.arange(0, self.dim, 2).float().to(device) / self.dim)
            )
            self.register_buffer("inv_freq", inv_freq, persistent=False)

        t = torch.arange(
            self.max_seq_len_cached, device=device, dtype=self.inv_freq.dtype
        )

        freqs = torch.outer(t, self.inv_freq)
        # Different from paper, but it uses a different permutation in order to obtain the same calculation
        emb = torch.cat((freqs, freqs), dim=-1)
        self.register_buffer("cos_cached", emb.cos().to(dtype), persistent=False)
        self.register_buffer("sin_cached", emb.sin().to(dtype), persistent=False)


# Inverse dim formula to find dim based on number of rotations
def yarn_find_correction_dim(
    num_rotations, dim, base=10000, max_position_embeddings=2048
):
    return (dim * math.log(max_position_embeddings / (num_rotations * 2 * math.pi))) / (
        2 * math.log(base)
    )


# Find dim range bounds based on rotations
def yarn_find_correction_range(
    low_rot, high_rot, dim, base=10000, max_position_embeddings=2048
):
    low = math.floor(
        yarn_find_correction_dim(low_rot, dim, base, max_position_embeddings)
    )
    high = math.ceil(
        yarn_find_correction_dim(high_rot, dim, base, max_position_embeddings)
    )
    return max(low, 0), min(high, dim - 1)  # Clamp values just in case


def yarn_get_mscale(scale=1, mscale=1):
    if scale <= 1:
        return 1.0
    return 0.1 * mscale * math.log(scale) + 1.0


def yarn_linear_ramp_mask(min, max, dim):
    if min == max:
        max += 0.001  # Prevent singularity

    linear_func = (torch.arange(dim, dtype=torch.float32) - min) / (max - min)
    ramp_func = torch.clamp(linear_func, 0, 1)
    return ramp_func


class DeepseekV2YarnRotaryEmbedding(DeepseekV2RotaryEmbedding):

    def __init__(
        self,
        config,
        dim,
        max_position_embeddings=2048,
        base=10000,
        device=None,
        scaling_factor=1.0,
        original_max_position_embeddings=4096,
        beta_fast=32,
        beta_slow=1,
        mscale=1,
        mscale_all_dim=0,
    ):
        self.scaling_factor = scaling_factor
        self.original_max_position_embeddings = original_max_position_embeddings
        self.beta_fast = beta_fast
        self.beta_slow = beta_slow
        self.mscale = mscale
        self.mscale_all_dim = mscale_all_dim
        super().__init__(config, dim, max_position_embeddings, base, device)

    def _set_cos_sin_cache(self, seq_len, device, dtype):
        self.max_seq_len_cached = seq_len
        dim = self.dim

        freq_extra = 1.0 / (
            self.base
            ** (torch.arange(0, dim, 2, dtype=torch.float32, device=device) / dim)
        )
        freq_inter = 1.0 / (
            self.scaling_factor
            * self.base
            ** (torch.arange(0, dim, 2, dtype=torch.float32, device=device) / dim)
        )

        low, high = yarn_find_correction_range(
            self.beta_fast,
            self.beta_slow,
            dim,
            self.base,
            self.original_max_position_embeddings,
        )
        inv_freq_mask = 1.0 - yarn_linear_ramp_mask(low, high, dim // 2).to(
            device=device, dtype=torch.float32
        )
        inv_freq = freq_inter * (1 - inv_freq_mask) + freq_extra * inv_freq_mask
        self.register_buffer("inv_freq", inv_freq, persistent=False)

        t = torch.arange(seq_len, device=device, dtype=torch.float32)

        freqs = torch.outer(t, inv_freq)

        _mscale = float(
            yarn_get_mscale(self.scaling_factor, self.mscale)
            / yarn_get_mscale(self.scaling_factor, self.mscale_all_dim)
        )

        emb = torch.cat((freqs, freqs), dim=-1)
        self.register_buffer(
            "cos_cached", (emb.cos() * _mscale).to(dtype), persistent=False
        )
        self.register_buffer(
            "sin_cached", (emb.sin() * _mscale).to(dtype), persistent=False
        )


# Copied from transformers.models.llama.modeling_llama.rotate_half
def rotate_half(x):
    """Rotates half the hidden dims of the input."""
    x1 = x[..., : x.shape[-1] // 2]
    x2 = x[..., x.shape[-1] // 2 :]
    return torch.cat((-x2, x1), dim=-1)


# Copied from transformers.models.llama.modeling_llama.apply_rotary_pos_emb
def apply_rotary_pos_emb(q, k, cos, sin, position_ids, unsqueeze_dim=1):
    """Applies Rotary Position Embedding to the query and key tensors.

    Args:
        q (`torch.Tensor`): The query tensor.
        k (`torch.Tensor`): The key tensor.
        cos (`torch.Tensor`): The cosine part of the rotary embedding.
        sin (`torch.Tensor`): The sine part of the rotary embedding.
        position_ids (`torch.Tensor`):
            The position indices of the tokens corresponding to the query and key tensors. For example, this can be
            used to pass offsetted position ids when working with a KV-cache.
        unsqueeze_dim (`int`, *optional*, defaults to 1):
            The 'unsqueeze_dim' argument specifies the dimension along which to unsqueeze cos[position_ids] and
            sin[position_ids] so that they can be properly broadcasted to the dimensions of q and k. For example, note
            that cos[position_ids] and sin[position_ids] have the shape [batch_size, seq_len, head_dim]. Then, if q and
            k have the shape [batch_size, heads, seq_len, head_dim], then setting unsqueeze_dim=1 makes
            cos[position_ids] and sin[position_ids] broadcastable to the shapes of q and k. Similarly, if q and k have
            the shape [batch_size, seq_len, heads, head_dim], then set unsqueeze_dim=2.
    Returns:
        `tuple(torch.Tensor)` comprising of the query and key tensors rotated using the Rotary Position Embedding.
    """
    cos = cos[position_ids].unsqueeze(unsqueeze_dim) # BSND->BNSD
    sin = sin[position_ids].unsqueeze(unsqueeze_dim) # BSND->BNSD

    b, h, s, d = q.shape
    q = q.view(b, h, s, d // 2, 2).transpose(4, 3).reshape(b, h, s, d)

    b, h, s, d = k.shape
    k = k.view(b, h, s, d // 2, 2).transpose(4, 3).reshape(b, h, s, d)

    q_embed = (q * cos) + (rotate_half(q) * sin)
    k_embed = (k * cos) + (rotate_half(k) * sin)
    return q_embed, k_embed


def rope_single(tensor, cos, sin):
    q_embed = torch_npu.npu_interleave_rope(tensor, cos, sin)
    return q_embed


def npu_apply_rotary_pos_emb(q, k, cos, sin, position_ids, layer_idx):
    if False:
        b, h, s, d = q.shape
        q = q.view(b, h, s, d // 2, 2).transpose(4, 3).reshape(b, h, s, d)
        q = q.transpose(1, 2) # BSND

        b, h, s, d = k.shape
        k = k.view(b, h, s, d // 2, 2).transpose(4, 3).reshape(b, h, s, d)
        k = k.transpose(1, 2) # BSND

        q_embed, k_embed = torch_npu.npu_apply_rotary_pos_emb(q, k, cos, sin, layout='BSH')
        
        q_embed = q_embed.transpose(1, 2)
        k_embed = k_embed.transpose(1, 2)
        
    else:
        q_embed = rope_single(q, cos, sin)
        k_embed = rope_single(k, cos, sin)
        
    return q_embed, k_embed



# Copied from transformers.models.llama.modeling_llama.repeat_kv
def repeat_kv(hidden_states: torch.Tensor, n_rep: int) -> torch.Tensor:
    """
    This is the equivalent of torch.repeat_interleave(x, dim=1, repeats=n_rep). The hidden states go from (batch,
    num_key_value_heads, seqlen, head_dim) to (batch, num_attention_heads, seqlen, head_dim)
    """
    batch, num_key_value_heads, slen, head_dim = hidden_states.shape
    if n_rep == 1:
        return hidden_states
    hidden_states = hidden_states[:, :, None, :, :].expand(
        batch, num_key_value_heads, n_rep, slen, head_dim
    )
    return hidden_states.reshape(batch, num_key_value_heads * n_rep, slen, head_dim)


def _init_rope(self):
    if self.config.rope_scaling is None:
        if "deepseek" in self.config.model_type:
            dim = self.config.qk_rope_head_dim 
        else:
            dim = self.config.attn_intermediate_size // self.config.num_attention_heads

        self.rotary_emb = DeepseekV2RotaryEmbedding(
            self.config,
            dim,
            max_position_embeddings=self.config.max_position_embeddings,
            base=self.config.rope_theta,
        )
    else:
        scaling_type = self.config.rope_scaling["type"]
        scaling_factor = self.config.rope_scaling["factor"]
        if scaling_type == "linear":
            self.rotary_emb = DeepseekV2LinearScalingRotaryEmbedding(
                self.config,
                self.config.qk_rope_head_dim,
                max_position_embeddings=self.config.max_position_embeddings,
                scaling_factor=scaling_factor,
                base=self.config.rope_theta,
            )
        elif scaling_type == "dynamic":
            self.rotary_emb = DeepseekV2DynamicNTKScalingRotaryEmbedding(
                self.config,
                self.config.qk_rope_head_dim,
                max_position_embeddings=self.config.max_position_embeddings,
                scaling_factor=scaling_factor,
                base=self.config.rope_theta,
            )
        elif scaling_type == "yarn":
            kwargs = {
                key: self.config.rope_scaling[key]
                for key in [
                    "original_max_position_embeddings",
                    "beta_fast",
                    "beta_slow",
                    "mscale",
                    "mscale_all_dim",
                ]
                if key in self.config.rope_scaling
            }
            self.rotary_emb = DeepseekV2YarnRotaryEmbedding(
                self.config,
                self.config.qk_rope_head_dim,
                max_position_embeddings=self.config.max_position_embeddings,
                scaling_factor=scaling_factor,
                base=self.config.rope_theta,
                **kwargs,
            )
        else:
            raise ValueError(f"Unknown RoPE scaling type {scaling_type}")

DeepseekV2_START_DOCSTRING = r"""
    This model inherits from [`PreTrainedModel`]. Check the superclass documentation for the generic methods the
    library implements for all its model (such as downloading or saving, resizing the input embeddings, pruning heads
    etc.)

    This model is also a PyTorch [torch.nn.Module](https://pytorch.org/docs/stable/nn.html#torch.nn.Module) subclass.
    Use it as a regular PyTorch Module and refer to the PyTorch documentation for all matter related to general usage
    and behavior.

    Parameters:
        config ([`DeepseekV2Config`]):
            Model configuration class with all the parameters of the model. Initializing with a config file does not
            load the weights associated with the model, only the configuration. Check out the
            [`~PreTrainedModel.from_pretrained`] method to load the model weights.
"""
@add_start_docstrings(
    "The bare DeepseekV2 Model outputting raw hidden-states without any specific head on top.",
    DeepseekV2_START_DOCSTRING,
)
class DeepseekV2PreTrainedModel(PreTrainedModel):
    config_class = DeepseekV2Config
    base_model_prefix = "model"
    supports_gradient_checkpointing = True
    _no_split_modules = ["DeepseekV2DecoderLayer"]
    _skip_keys_device_placement = "past_key_values"
    _supports_flash_attn_2 = True
    _supports_cache_class = True

    def _init_weights(self, module):
        pass
        # std = self.config.initializer_range
        # if isinstance(module, nn.Linear):
        #     module.weight.data.normal_(mean=0.0, std=std)
        #     if module.bias is not None:
        #         module.bias.data.zero_()
        # elif isinstance(module, nn.Embedding):
        #     module.weight.data.normal_(mean=0.0, std=std)
        #     if module.padding_idx is not None:
        #         module.weight.data[module.padding_idx].zero_()


def init_comm_group(
    global_rank,
    group_num, 
    world_size,
    group_stride=1,
    group_name="default",
    return_name=False,
    rank_offset=0
):
    # 预期创建 group_num 个通信域, 每个通信域的大小是 group_size
    ## 默认 group_num * group_size == world_size
    # 每个通信域中, rank的跨度是group_stride
    assert world_size is not None and world_size > 0
    group_size = world_size // group_num

    cur_group_set = None
    for group_id in range(group_num):
        # 构建每个通信域的首rank
        if group_stride == 1:
            # 连续通信域场景，起始地址为 0, 4, 8, 12。每个通信域内为 [0, 1, 2, 3]
            start_rank_id = group_id * group_size + rank_offset
            init_rank_id = global_rank // group_size * group_size
        else:
            # 跳跃通信域场景，起始地址为 0, 1, 2, 3。每个通信域内为 [0, 4, 8, 12]
            start_rank_id = group_id + rank_offset
            init_rank_id = global_rank % group_num

        cur_group_list = [start_rank_id + i * group_stride for i in range(group_size)]
        if _world._default_pg is not None:
            cur_group = dist.new_group(cur_group_list)
        else:
            cur_group = None
        print(f"group_name:{group_name} group_id:{group_id} group_num:{group_num} "
              f"start_rank_id:{start_rank_id} init_rank_id:{init_rank_id}")
        if start_rank_id == init_rank_id:
            cur_group_set = cur_group
            print(f"group_name is {group_name}, group_list: {cur_group_list}, "
                  f"group_id:{group_id}, group_num:{group_num}", flush=True)
    if (not return_name) or save_model_switch:
        print(f"init_comm_group 1 group_name:{group_name} cur_group_set:{cur_group_set}")
        return cur_group_set
    else:
        print(f"init_comm_group 2 group_name:{group_name}")
        return cur_group_set._get_backend(torch.device("npu")).get_hccl_comm_name(global_rank)


def half_batch(ori_tensor, split_dim):
    if ori_tensor is None:
        return (None, None)
    tensor_set = torch.chunk(ori_tensor, 2, dim=split_dim)
    return (tensor_set[0], tensor_set[1])


def one_third_batch(ori_tensor, split_dim):
    if ori_tensor is None:
        return (None, None, None)
    tensor_set = torch.chunk(ori_tensor, 3, dim=split_dim)
    return (tensor_set[0], tensor_set[1], tensor_set[2])


class NpuLinear(nn.Module):
    def __init__(self, in_feature, out_feature, bias: bool = False):
        super().__init__()
        fix_rand_seed()
        self.weight = nn.Parameter(torch.rand((out_feature, in_feature), dtype=model_dtype), requires_grad=False)
        self.bias = None
        if bias is not None and bias:
            self.bias = nn.Parameter(torch.rand((out_feature,), dtype=model_dtype))

    def forward(self, x):
        """
        x and weight should be 2-D tensors both.
        weight should be transposed in cast_format before do matmul.
        """
        origin_shape = x.size()
        x = x.view(-1, origin_shape[-1])
        out = torch.matmul(x, self.weight.data)
        out = out.view(*origin_shape[:-1], -1)
        return out


class FakeContextManager:
    def __init__(self) -> None:
        pass
    def __enter__(self):
        pass
    def __exit__(self, type, value, traceback):
        pass


def NpuStreamSwitch(open: bool, stream_tag: str, stream_priority: int = 0):
    if open:
        return tng.scope.npu_stream_switch(stream_tag, stream_priority)
    else:
        return FakeContextManager()


def SuperKernelScope(open: bool, scope: str, options: str = None):
    if open:
        return tng.scope.super_kernel(scope, options)
    else:
        return FakeContextManager()

def NpuLimitCoreNum(open: bool, op_aicore_num: int, op_vectorcore_num: int):
    if open:
        return tng.scope.limit_core_num(op_aicore_num, op_vectorcore_num)
    else:
        return FakeContextManager()


# max_len须小于_PAGE_ATTENTION_SETTING["max_length"]
def get_actual_seq_len_list(filename, cnt, row_offset=1, max_len=actual_seq_len - max_new_tokens):
    from openpyxl import load_workbook
    wb = load_workbook(filename, data_only=True)  # 加载 Excel 文件
    ws = wb.active  # 默认读取第一个工作表

    result = []
    for row in range(row_offset, ws.max_row + 1):
        rel_seq_len = ws[f'A{row}'].value
        if isinstance(rel_seq_len, int) and 1 < rel_seq_len <= max_len:
            result.append(rel_seq_len)
            if len(result) == cnt:
                break
    return result


def to_transpose_nz(tensor, transpose_contiguous=False):
    if transpose_contiguous:
        tensor.data = tensor.data.transpose(-2, -1).contiguous()
    return torch_npu.npu_format_cast(tensor.data, 29)
